import os
from docx import Document
from openpyxl import load_workbook
import PyPDF2

# Extrae metadatos de un archivo .docx
def extraer_metadatos_docx(archivo):
    doc = Document(archivo)
    propiedades = doc.core_properties
    return {
        'título': propiedades.title,
        'autor': propiedades.author,
        'fecha de creación': propiedades.created,
        'última modificación': propiedades.modified,
    }

# Extrae metadatos de un archivo .xlsx
def extraer_metadatos_xlsx(archivo):
    wb = load_workbook(filename=archivo, data_only=True)
    propiedades = wb.properties
    return {
        'título': propiedades.title,
        'autor': propiedades.creator,
        'fecha de creación': propiedades.created,
        'última modificación': propiedades.modified,
    }

# Extrae metadatos de un archivo .pdf
def extraer_metadatos_pdf(archivo):
    with open(archivo, 'rb') as file:
        reader = PyPDF2.PdfReader(file)
        info = reader.metadata
        return {
            'autor': info.get('/Author'),
            'creador': info.get('/Creator'),
            'productor': info.get('/Producer'),
            'título': info.get('/Title'),
            'fecha de creación': info.get('/CreationDate'),
            'última modificación': info.get('/ModDate'),
        }

# Muestra los metadatos de todos los archivos en la ruta especificada
def mostrar_metadatos(ruta):
    for archivo in os.listdir(ruta):
        ruta_completa = os.path.join(ruta, archivo)
        if archivo.endswith('.docx'):
            metadatos = extraer_metadatos_docx(ruta_completa)
            tipo = 'DOCX'
        elif archivo.endswith('.xlsx'):
            metadatos = extraer_metadatos_xlsx(ruta_completa)
            tipo = 'XLSX'
        elif archivo.endswith('.pdf'):
            metadatos = extraer_metadatos_pdf(ruta_completa)
            tipo = 'PDF'
        else:
            continue

        print(f'Metadatos para {tipo} archivo {archivo}:')
        for key, value in metadatos.items():
            print(f'{key}: {value}')
        print('----------------------------------------')

# Define la ruta del directorio y muestra los metadatos de los archivos
ruta_directorio = '/home/kali/Documents/'
mostrar_metadatos(ruta_directorio)
